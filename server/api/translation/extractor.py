from pptx import Presentation

from backend.config import ENGLISH_PREFIX


# ============================================================
# EXTRACT ONE SLIDE
# ============================================================

def extract_slide_text(slide):

    elements = []

    for shape in slide.shapes:

        # ----------------------------------------------------
        # Only process text boxes
        # ----------------------------------------------------

        if not shape.has_text_frame:
            continue

        shape_name = shape.name

        # ----------------------------------------------------
        # Only process EN_ shapes
        # ----------------------------------------------------

        if not shape_name.startswith(
            ENGLISH_PREFIX
        ):
            continue

        paragraphs = []

        # ----------------------------------------------------
        # Extract every paragraph separately
        # ----------------------------------------------------

        for paragraph in shape.text_frame.paragraphs:

            text = paragraph.text.strip()

            if not text:
                continue

            paragraphs.append({

                "text": text

            })

        # ----------------------------------------------------
        # Ignore empty text boxes
        # ----------------------------------------------------

        if not paragraphs:
            continue

        elements.append({

            "shape_name": shape_name,

            "paragraphs": paragraphs

        })

    return elements


# ============================================================
# EXTRACT ENTIRE PRESENTATION
# ============================================================

def extract_presentation(ppt_path):

    prs = Presentation(ppt_path)

    presentation_data = []

    for slide_number, slide in enumerate(
        prs.slides,
        start=1
    ):

        print(
            f"Reading slide "
            f"{slide_number}..."
        )

        slide_data = {

            "slide_number":
                slide_number,

            "elements":
                extract_slide_text(slide)

        }

        presentation_data.append(
            slide_data
        )

    return presentation_data