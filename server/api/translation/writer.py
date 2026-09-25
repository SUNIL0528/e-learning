from pathlib import Path

from pptx import Presentation

from backend.config import (
    ENGLISH_PREFIX,
    TRANSLATION_PREFIX
)

from .layout import (
    insert_translation
)


# ============================================================
# GET TRANSLATION BOX NAME
# ============================================================

def get_translation_shape_name(
    english_shape_name
):

    if not english_shape_name.startswith(
        ENGLISH_PREFIX
    ):
        return None

    suffix = english_shape_name[
        len(ENGLISH_PREFIX):
    ]

    return (
        TRANSLATION_PREFIX +
        suffix
    )


# ============================================================
# FIND SHAPE
# ============================================================

def find_shape(
    slide,
    shape_name
):

    for shape in slide.shapes:

        if shape.name == shape_name:

            return shape

    return None


# ============================================================
# WRITE ONE SLIDE
# ============================================================

def write_slide_translations(
    slide,
    slide_translation
):

    for element in slide_translation[
        "elements"
    ]:

        english_shape_name = (
            element[
                "shape_name"
            ]
        )

        translated_paragraphs = (
            element[
                "paragraphs"
            ]
        )

        # ----------------------------------------------------
        # EN_title -> TR_title
        # EN_description -> TR_description
        # ----------------------------------------------------

        target_shape_name = (
            get_translation_shape_name(
                english_shape_name
            )
        )

        if not target_shape_name:
            continue

        # ----------------------------------------------------
        # Find English shape
        # ----------------------------------------------------

        source_shape = find_shape(
            slide,
            english_shape_name
        )

        if source_shape is None:

            print(
                f"WARNING: English shape "
                f"'{english_shape_name}' "
                f"not found."
            )

            continue

        # ----------------------------------------------------
        # Find translation shape
        # ----------------------------------------------------

        target_shape = find_shape(
            slide,
            target_shape_name
        )

        if target_shape is None:

            print(
                f"WARNING: Translation shape "
                f"'{target_shape_name}' "
                f"not found."
            )

            continue

        # ----------------------------------------------------
        # Write translation
        # ----------------------------------------------------

        insert_translation(
            source_shape,
            target_shape,
            translated_paragraphs
        )


# ============================================================
# CREATE TRANSLATED PPT
# ============================================================

def create_translated_ppt(
    input_ppt,
    output_ppt,
    translated_presentation
):

    input_ppt = Path(
        input_ppt
    )

    output_ppt = Path(
        output_ppt
    )

    # --------------------------------------------------------
    # Create output directory
    # --------------------------------------------------------

    output_ppt.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    print(
        "\nOpening original PPT..."
    )

    print(
        f"Input : {input_ppt}"
    )

    print(
        f"Output: {output_ppt}"
    )

    # --------------------------------------------------------
    # IMPORTANT:
    #
    # Presentation() loads the original PPT
    # into memory.
    #
    # prs.save(output_ppt) writes a NEW file.
    #
    # The input PPT is NOT modified.
    # --------------------------------------------------------

    prs = Presentation(
        str(input_ppt)
    )

    total_slides = len(
        translated_presentation
    )

    # --------------------------------------------------------
    # Process slides
    # --------------------------------------------------------

    for slide_index, slide_translation in enumerate(
        translated_presentation
    ):

        print(
            f"Writing slide "
            f"{slide_index + 1}/"
            f"{total_slides}"
        )

        slide = prs.slides[
            slide_index
        ]

        write_slide_translations(
            slide,
            slide_translation
        )

    # --------------------------------------------------------
    # Save NEW translated PPT
    # --------------------------------------------------------

    prs.save(
        str(output_ppt)
    )

    print(
        "\nTranslated PPT created successfully."
    )

    print(
        f"Output: {output_ppt}"
    )

    print(
        f"Original preserved: {input_ppt}"
    )